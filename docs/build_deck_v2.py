from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ───────────────────────────────────────────
C_DARK    = RGBColor(0x0A, 0x0A, 0x0F)
C_PRIMARY = RGBColor(0x63, 0x66, 0xF1)
C_ACCENT  = RGBColor(0xA7, 0x8B, 0xFA)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_MUTED   = RGBColor(0x94, 0xA3, 0xB8)
C_CARD    = RGBColor(0x12, 0x12, 0x1E)
C_SOFT    = RGBColor(0x1E, 0x1E, 0x30)
C_GREEN   = RGBColor(0x34, 0xD3, 0x99)
C_AMBER   = RGBColor(0xFB, 0xBF, 0x24)
C_ROSE    = RGBColor(0xFB, 0x71, 0x85)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]

# ── Helpers ───────────────────────────────────────────

def slide():
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C_DARK
    return s

def box(s, text, x, y, w, h,
        size=20, bold=False, color=C_WHITE,
        align=PP_ALIGN.LEFT, wrap=True):
    tf = s.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size  = Pt(size)
    r.font.bold  = bold
    r.font.color.rgb = color
    r.font.name  = "Calibri"
    return tf

def rect(s, x, y, w, h, fill=C_CARD):
    sh = s.shapes.add_shape(1, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh

def watermark(s):
    box(s, "Whiter.ai", W-Inches(1.6), H-Inches(.48), Inches(1.5), Inches(.38),
        size=11, bold=True, color=RGBColor(0x22,0x22,0x32), align=PP_ALIGN.RIGHT)

def counter(s, n, total):
    box(s, f"{n} / {total}", Inches(.5), H-Inches(.48), Inches(1), Inches(.38),
        size=10, color=RGBColor(0x30,0x30,0x45))

def eyebrow(s, text):
    r = rect(s, Inches(.75), Inches(.52), Inches(2.4), Inches(.34), fill=RGBColor(0x16,0x16,0x28))
    tf = r.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text.upper()
    run.font.size  = Pt(9)
    run.font.bold  = True
    run.font.color.rgb = C_ACCENT
    run.font.name  = "Calibri"

def note(s, text):
    s.notes_slide.notes_text_frame.text = text

def pill(s, label, x, y, color=C_PRIMARY):
    w = Inches(2.0); h = Inches(.38)
    r = rect(s, x, y, w, h, fill=RGBColor(0x16,0x16,0x2A))
    tf = r.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size  = Pt(9)
    run.font.bold  = True
    run.font.color.rgb = color
    run.font.name  = "Calibri"

TOTAL = 7

# ═══════════════════════════════════════════════════════
# 1. HOOK
# ═══════════════════════════════════════════════════════
s = slide()
rect(s, W-Inches(5.5), Inches(-1), Inches(7), Inches(9.5), fill=RGBColor(0x0C,0x0C,0x20))

box(s, "Did you know…", Inches(.75), Inches(1.1), Inches(11), Inches(.7),
    size=20, color=C_ACCENT)

box(s, "your audience forgets\n70% of what you just said\nwithin 24 hours?",
    Inches(.75), Inches(1.85), Inches(10), Inches(3.2),
    size=52, bold=True, color=C_WHITE)

box(s, "Not because they weren't paying attention.\nBecause nothing kept them there.",
    Inches(.75), Inches(5.2), Inches(7.5), Inches(.9),
    size=18, color=C_MUTED)

watermark(s); counter(s, 1, TOTAL)
note(s, "Pause after the stat. Let it land. Then: 'We built something about that.'")


# ═══════════════════════════════════════════════════════
# 2. PROBLEM + CUSTOMER
# ═══════════════════════════════════════════════════════
s = slide()
eyebrow(s, "The Problem")

box(s, "Today, professionals who present live\nstruggle to keep their audience engaged\nbecause their slides are built before\nthe conversation happens.",
    Inches(.75), Inches(1.1), Inches(8.5), Inches(2.8),
    size=28, bold=True, color=C_WHITE)

CW = Inches(3.8); CH = Inches(2.0); Y = Inches(4.2); GAP = Inches(.22)

for i, (label, lc, title, body) in enumerate([
    ("WHO",               C_ACCENT,  "Corporate teams, SMEs\n& live speakers",
     "Anyone who presents live without a design team behind them."),
    ("THE PAIN",          C_ROSE,    "Hours of prep.\nZero retention.",
     "Slides built in advance disconnect from the live conversation. Audiences switch off."),
    ("WHY SOLUTIONS FAIL", C_AMBER,  "PowerPoint is for\ndesign, not delivery.",
     "AI notetakers arrive post-meeting. Nothing generates visuals live, as you speak."),
]):
    x = Inches(.75) + i * (CW + GAP)
    rect(s, x, Y, CW, CH, fill=C_CARD)
    box(s, label, x+Inches(.2), Y+Inches(.15), CW-Inches(.4), Inches(.3),
        size=9, bold=True, color=lc)
    box(s, title, x+Inches(.2), Y+Inches(.48), CW-Inches(.4), Inches(.6),
        size=14, bold=True, color=C_WHITE)
    box(s, body, x+Inches(.2), Y+Inches(1.12), CW-Inches(.4), CH-Inches(1.2),
        size=10, color=C_MUTED)

watermark(s); counter(s, 2, TOTAL)
note(s, "The three boxes = who, pain, why nothing fixes it. Land each one.")


# ═══════════════════════════════════════════════════════
# 3. SOLUTION STATEMENT
# ═══════════════════════════════════════════════════════
s = slide()
eyebrow(s, "The Solution")

rect(s, Inches(-0.5), Inches(2.4), W+Inches(1), Inches(2.8), fill=RGBColor(0x0D,0x0D,0x22))

box(s, "We built", Inches(.75), Inches(1.1), Inches(10), Inches(.7),
    size=26, color=C_MUTED)

box(s, "Whiter.ai", Inches(.75), Inches(1.75), Inches(6), Inches(.95),
    size=54, bold=True, color=C_PRIMARY)

box(s, "The AI-powered whiteboard that listens to your speech and generates\nsynchronized visual slides for every attendee — in real time.",
    Inches(.75), Inches(2.75), Inches(11.2), Inches(1.1),
    size=20, color=C_WHITE)

box(s, "No prep.    No design.    No disconnect.    Just speak.",
    Inches(.75), Inches(3.95), Inches(11), Inches(.55),
    size=16, bold=True, color=C_ACCENT)

# Three outcomes
for i, (val, desc) in enumerate([
    ("Real-time", "Slides appear as you speak"),
    ("Any device", "Audience joins via URL — no install"),
    ("Any language", "8 languages supported"),
]):
    x = Inches(.75) + i * Inches(4.1)
    rect(s, x, Inches(5.0), Inches(3.8), Inches(1.4), fill=C_CARD)
    box(s, val, x+Inches(.2), Inches(5.1), Inches(3.5), Inches(.45),
        size=16, bold=True, color=C_WHITE)
    box(s, desc, x+Inches(.2), Inches(5.55), Inches(3.5), Inches(.6),
        size=11, color=C_MUTED)

watermark(s); counter(s, 3, TOTAL)
note(s, "Read the solution statement slowly. Then let the three boxes sink in.")


# ═══════════════════════════════════════════════════════
# 4. MARKET
# ═══════════════════════════════════════════════════════
s = slide()
eyebrow(s, "Market Opportunity")

box(s, "A $30B market\nstuck in 1987.",
    Inches(.75), Inches(1.1), Inches(9), Inches(1.8),
    size=42, bold=True, color=C_WHITE)

box(s, "PowerPoint launched in 1987. Nothing has changed the live presentation experience since.",
    Inches(.75), Inches(2.95), Inches(8), Inches(.6),
    size=14, color=C_MUTED)

# TAM / SAM / SOM funnel
for i, (label, val, desc, col) in enumerate([
    ("TAM", "$30B",  "Global presentation & meeting productivity software", C_PRIMARY),
    ("SAM", "$8B",   "Professionals who present live regularly\n(corporate, SME, speakers)", C_ACCENT),
    ("SOM", "€1M",   "Year 1 — European beachhead:\nSME sales teams, consultants & independent speakers", C_GREEN),
]):
    W2 = Inches(3.8) - i * Inches(.4)
    x  = Inches(.75) + i * Inches(.2)
    y  = Inches(3.8) + i * Inches(1.0)
    rect(s, x, y, W2, Inches(.9), fill=C_CARD)
    box(s, label, x+Inches(.2), y+Inches(.1), Inches(.7), Inches(.35),
        size=9, bold=True, color=col)
    box(s, val, x+Inches(.95), y+Inches(.08), Inches(1.2), Inches(.4),
        size=20, bold=True, color=C_WHITE)
    box(s, desc, x+Inches(2.25), y+Inches(.1), W2-Inches(2.45), Inches(.7),
        size=9, color=C_MUTED)

watermark(s); counter(s, 4, TOTAL)
note(s, "TAM = total addressable. SAM = who we can realistically sell to. SOM = year 1 target. Keep it brief.")


# ═══════════════════════════════════════════════════════
# 5. HOW IT WORKS + REVENUE
# ═══════════════════════════════════════════════════════
s = slide()
eyebrow(s, "How It Works")

# Steps
for i, (num, title, body) in enumerate([
    ("1", "Set up",  "Choose your language, visual theme, and give the AI a quick briefing on the session."),
    ("2", "Speak",   "Whiter.ai transcribes in real time. AI extracts key insights and generates slides instantly."),
    ("3", "Share",   "Audience opens a URL — slides appear on every screen as you talk. QR summary at the end."),
]):
    y = Inches(1.3) + i * Inches(1.4)
    rect(s, Inches(.75), y+Inches(.1), Inches(.55), Inches(.55), fill=C_PRIMARY)
    box(s, num, Inches(.75), y, Inches(.55), Inches(.65),
        size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    box(s, title, Inches(1.45), y, Inches(4.0), Inches(.4),
        size=14, bold=True, color=C_WHITE)
    box(s, body,  Inches(1.45), y+Inches(.4), Inches(5.8), Inches(.85),
        size=11, color=C_MUTED)

# Revenue — 3 packages
box(s, "Three plans. Every use case.", Inches(7.2), Inches(1.1), Inches(5.6), Inches(.5),
    size=14, bold=True, color=C_WHITE)

plans = [
    ("Speaker",  "For coaches, trainers\n& independent speakers", C_ACCENT),
    ("SME",      "For small teams\n& founders", C_PRIMARY),
    ("Business", "Custom branding,\nintegrations & API", C_GREEN),
]
PW = Inches(1.7); PH = Inches(4.2); PY = Inches(1.75)
for i, (name, desc, col) in enumerate(plans):
    x = Inches(7.2) + i * (PW + Inches(.2))
    rect(s, x, PY, PW, PH, fill=C_CARD)
    box(s, name, x+Inches(.15), PY+Inches(.2), PW-Inches(.3), Inches(.38),
        size=13, bold=True, color=col)
    box(s, desc, x+Inches(.15), PY+Inches(.65), PW-Inches(.3), PH-Inches(.85),
        size=9, color=C_MUTED)

watermark(s); counter(s, 5, TOTAL)
note(s, "Quick walkthrough of the 3 steps, then land the 3 packages. Don't linger.")


# ═══════════════════════════════════════════════════════
# 6. DEMO
# ═══════════════════════════════════════════════════════
s = slide()

# Glow
rect(s, Inches(2), Inches(.5), Inches(9), Inches(6.5), fill=RGBColor(0x0C,0x0C,0x22))

box(s, "Live Demo", Inches(.75), Inches(1.3), Inches(11.8), Inches(.6),
    size=16, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)

box(s, "This is running right now.",
    Inches(.75), Inches(2.0), Inches(11.8), Inches(1.1),
    size=52, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

box(s, "Open the URL on screen. Watch what happens as I keep talking.",
    Inches(.75), Inches(3.2), Inches(11.8), Inches(.6),
    size=16, color=C_MUTED, align=PP_ALIGN.CENTER)

# URL box
rect(s, Inches(4.0), Inches(4.1), Inches(5.3), Inches(1.2), fill=C_CARD)
box(s, "⬤  Live session active",
    Inches(4.0), Inches(4.2), Inches(5.3), Inches(.35),
    size=10, color=RGBColor(0x47,0x55,0x69), align=PP_ALIGN.CENTER)
box(s, "← paste board URL here before going on stage",
    Inches(4.0), Inches(4.55), Inches(5.3), Inches(.55),
    size=12, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)

watermark(s); counter(s, 6, TOTAL)
note(s, "DEMO MOMENT — start Whiter.ai BEFORE stepping up. Replace text above with real ?board= URL. Pause 5 seconds after showing it. Say nothing. Let the room open it on their phones.")


# ═══════════════════════════════════════════════════════
# 7. CLOSE
# ═══════════════════════════════════════════════════════
s = slide()

rect(s, Inches(-0.5), Inches(2.2), W+Inches(1), Inches(3.1), fill=RGBColor(0x0D,0x0D,0x22))

box(s, "The future meeting", Inches(.75), Inches(1.2), Inches(11), Inches(.9),
    size=46, bold=True, color=C_WHITE)
box(s, "has no deck.", Inches(.75), Inches(2.05), Inches(11), Inches(.9),
    size=46, bold=True, color=C_WHITE)
box(s, "Just conversation.", Inches(.75), Inches(2.9), Inches(11), Inches(.9),
    size=46, bold=True, color=C_PRIMARY)

box(s, "whiter.ai — try it, join the waitlist, or talk to us after this.",
    Inches(.75), Inches(4.15), Inches(9), Inches(.55),
    size=15, color=C_MUTED)

# Built stat
rect(s, Inches(.75), Inches(5.1), Inches(4.5), Inches(.9), fill=C_CARD)
box(s, "Built in 48 hours at ESMT Berlin Sustainability Bootcamp 2026",
    Inches(.95), Inches(5.18), Inches(4.1), Inches(.72),
    size=10, color=C_MUTED)

watermark(s); counter(s, 7, TOTAL)
note(s, "Close slowly. Pause on 'Just conversation.' Then smile. Done. No ask, no begging — let them come to you.")


# ── Save ──────────────────────────────────────────────
out = "/Users/gautierserres/Downloads/launchpad-main/docs/Whiter_ai_Pitch_v2.pptx"
prs.save(out)
print(f"Saved: {out}")
